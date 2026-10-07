#!/usr/bin/env python3
"""Generate a completed STROBE checklist (cohort/cross-sectional items as
applicable to an ecological specialty-level panel) for the JHSRP submission."""

import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

BASE_DIR = os.path.dirname(os.path.abspath(__file__))


def build():
    doc = Document()

    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(10)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('STROBE Statement \u2014 checklist of items for observational '
                    'studies (ecological specialty-level panel)')
    run.bold = True
    run.font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Malpractice litigation, physician supply and service capacity '
                  'in Japan, 2008\u20132024: a national equivalence analysis')
    r.italic = True
    r.font.size = Pt(10)

    doc.add_paragraph()

    items = [
        ('Title and abstract', '1',
         '(a) Indicate the study\u2019s design with a commonly used term in the title or abstract\n'
         '(b) Provide in the abstract an informative and balanced summary of what was done and what was found',
         '(a) Title ("a national equivalence analysis"); Abstract (Observational specialty-level panel)\n'
         '(b) Structured abstract (Objectives, Methods, Results, Conclusions)'),

        ('Introduction', '', '', ''),
        ('Background/rationale', '2',
         'Explain the scientific background and rationale for the investigation being reported',
         'Introduction, paragraphs 1\u20135'),
        ('Objectives', '3',
         'State specific objectives, including any prespecified hypotheses',
         'Introduction, final paragraph'),

        ('Methods', '', '', ''),
        ('Study design', '4',
         'Present key elements of study design early in the paper',
         'Methods: Data sources, first paragraph (ecological specialty-level panel, 12 specialties, 2008\u20132024, nine biennial waves)'),
        ('Setting', '5',
         'Describe the setting, locations, and relevant dates, including periods of data collection',
         'Methods: Data sources (Japan, national administrative series, 2008\u20132024)'),
        ('Participants', '6',
         'Give the eligibility criteria, and the sources and methods of selection of participants',
         'Methods: Data sources (12 core clinical specialties for which the Supreme Court reports specialty-specific litigation); Supplementary Table 1'),
        ('Variables', '7',
         'Clearly define all outcomes, exposures, predictors, potential confounders, and effect modifiers',
         'Methods: Data sources and Statistical analysis (litigation rate exposure; physician and hospital growth outcomes; JOCS-CP indicator)'),
        ('Data sources/measurement', '8',
         'For each variable of interest, give sources of data and details of methods of assessment',
         'Methods: Data sources (MHLW Statistics of Physicians; Supreme Court closed claims; Survey of Medical Institutions; JMSR; Nikkei Telecom); Supplementary Table 1'),
        ('Bias', '9',
         'Describe any efforts to address potential sources of bias',
         'Methods (rate-based exposure to remove specialty-size confounding; measured waves only; cluster-robust inference); Results (confounder-adjusted sensitivity models); Limitations'),
        ('Study size', '10',
         'Explain how the study size was arrived at',
         'Methods: Data sources (all specialties with specialty-specific litigation statistics; all measured waves used)'),
        ('Quantitative variables', '11',
         'Explain how quantitative variables were handled in the analyses',
         'Methods: Statistical analysis (biennial log-change; standardised litigation rate; equivalence margins)'),
        ('Statistical methods', '12',
         '(a) Describe all statistical methods\n'
         '(b) Describe any methods used to examine subgroups and interactions\n'
         '(c) Explain how missing data were addressed\n'
         '(d) If applicable, describe analytical methods taking account of sampling strategy\n'
         '(e) Describe any sensitivity analyses',
         '(a) Methods: Statistical analysis (fixed-effects panel, clustered SEs, TOST equivalence testing)\n'
         '(b) No subgroup analyses\n'
         '(c) No missing data in the aggregate series analysed\n'
         '(d) Not applicable (population-level aggregate data)\n'
         '(e) Methods: Statistical analysis (annual hospital series, interpolated physician series, counts contrast, JMSR-adjusted and media-adjusted models)'),

        ('Results', '', '', ''),
        ('Participants', '13',
         '(a) Report numbers of individuals at each stage of study\n'
         '(b) Give reasons for non-participation at each stage\n'
         '(c) Consider use of a flow diagram',
         '(a) 12 specialties, 96 specialty-wave observations (aggregate ecological data; no individual participants)\n'
         '(b) Not applicable\n'
         '(c) Not applicable'),
        ('Descriptive data', '14',
         '(a) Give characteristics of study participants and information on exposures and potential confounders\n'
         '(b) Indicate the number of participants with missing data for each variable of interest',
         '(a) Results: Workforce and litigation trends; Table 1; Supplementary Figures 1\u20132\n'
         '(b) No missing data in the analysed series'),
        ('Outcome data', '15',
         'Report numbers of outcome events or summary measures',
         'Results (biennial log-change in physicians and hospitals; Table 1)'),
        ('Main results', '16',
         '(a) Give unadjusted estimates and, if applicable, confounder-adjusted estimates\n'
         '(b) Report category boundaries when continuous variables were categorised\n'
         '(c) If relevant, consider translating estimates of relative risk into absolute risk',
         '(a) Results: Primary association and equivalence; Table 2 (coefficients, 95% CIs, p values, n)\n'
         '(b) Continuous variables not categorised\n'
         '(c) Effects expressed as percentage biennial workforce change per 1-SD litigation rate'),
        ('Other analyses', '17',
         'Report other analyses done\u2014e.g., analyses of subgroups and interactions, and sensitivity analyses',
         'Results: sensitivity models (counts contrast, annual and interpolated series, JMSR and media adjustment, reverse specification); Supplementary Tables 2\u20134'),

        ('Discussion', '', '', ''),
        ('Key results', '18',
         'Summarise key results with reference to study objectives',
         'Discussion, opening paragraphs'),
        ('Limitations', '19',
         'Discuss limitations of the study, taking into account sources of potential bias or imprecision',
         'Discussion: Limitations (ecological design, 12 clusters, equivalence-margin judgement, exposure proxy, Japan context)'),
        ('Interpretation', '20',
         'Give a cautious overall interpretation of results considering objectives, limitations, '
         'multiplicity of analyses, results from similar studies, and other relevant evidence',
         'Discussion (equivalence interpretation, international literature, exploratory JOCS-CP association)'),
        ('Generalisability', '21',
         'Discuss the generalisability (external validity) of the study results',
         'Discussion (Japan-specific setting; transferable research question; explicit limits to generalisation)'),

        ('Other information', '', '', ''),
        ('Funding', '22',
         'Give the source of funding and the role of the funders for the present study',
         'Title page: Statements and Declarations \u2013 Funding (no financial support received)'),
    ]

    headers = ['Section/Topic', 'Item No.', 'Recommendation', 'Reported on page/section']
    t = doc.add_table(rows=1, cols=4)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, width in enumerate([Cm(3.0), Cm(1.2), Cm(8.0), Cm(6.0)]):
        t.columns[i].width = width

    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.name = 'Times New Roman'
        run.font.size = Pt(9)

    for topic, num, recommendation, location in items:
        row = t.add_row()
        cells = row.cells
        if not num and not recommendation:
            cells[0].text = ''
            p = cells[0].paragraphs[0]
            run = p.add_run(topic)
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(9)
            for j in range(1, 4):
                cells[j].text = ''
        else:
            for j, val in enumerate([topic, num, recommendation, location]):
                cells[j].text = ''
                p = cells[j].paragraphs[0]
                run = p.add_run(val)
                run.font.name = 'Times New Roman'
                run.font.size = Pt(9)

    out_path = os.path.join(BASE_DIR, 'JHSRP_STROBE_CHECKLIST.docx')
    doc.save(out_path)
    print(f'STROBE checklist saved to {out_path}')


if __name__ == '__main__':
    build()
